import json
import re
from pathlib import Path
from xml.sax.saxutils import escape

from ..ai import chat
from .local_tools import WORKSPACE

WORK_ROOT = Path(WORKSPACE) / "jarvis_projects"
WORK_ROOT.mkdir(parents=True, exist_ok=True)

def slug(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9_-]+", "_", value).strip("_")[:70] or "jarvis_project"

async def make_plan(request: str) -> dict:
    prompt = """Design a practical ECE engineering project.
Return JSON only with name, objective, novelty, components, blocks, connections,
firmware_modules, validation_steps, pcb, ppt_outline and report_sections.
Use components [{name,value_or_part,qty}], blocks [{name,x,y,width,height}],
connections [{from,to,signal}], and pcb {width_mm,height_mm,layers,notes}.
Mark uncertain hardware details as assumptions."""
    raw = await chat([
        {"role": "system", "content": prompt},
        {"role": "user", "content": request},
    ], temperature=0.2)
    try:
        return json.loads(raw)
    except Exception:
        match = re.search(r"\{.*\}", raw, re.S)
        if match:
            return json.loads(match.group(0))
        raise RuntimeError("JARVIS planner returned invalid JSON")

def write_docx(path: Path, plan: dict):
    from docx import Document
    doc = Document()
    doc.add_heading(plan["name"], 0)
    doc.add_paragraph(plan.get("objective", ""))
    for title, key in [
        ("Concept", "novelty"), ("Components", "components"),
        ("Architecture", "blocks"), ("Connections", "connections"),
        ("Firmware", "firmware_modules"), ("Validation", "validation_steps")
    ]:
        doc.add_heading(title, 1)
        values = plan.get(key, [])
        if isinstance(values, str):
            values = [values]
        for value in values:
            text = json.dumps(value, ensure_ascii=False) if isinstance(value, dict) else str(value)
            doc.add_paragraph(text, style="List Bullet")
    doc.save(path)

def write_pptx(path: Path, plan: dict):
    from pptx import Presentation
    prs = Presentation()
    outline = plan.get("ppt_outline") or [
        "Problem", "Proposed Solution", "Architecture", "Hardware",
        "Firmware", "PCB", "Testing", "Results", "Future Work"
    ]
    for index, title in enumerate(outline):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = str(title)
        body = slide.placeholders[1]
        if index == 0:
            value = plan.get("objective", "")
        elif index == 2:
            value = " -> ".join(str(x.get("name", "")) for x in plan.get("blocks", []) if isinstance(x, dict))
        elif index == 3:
            value = "\n".join(f'{x.get("name")} — {x.get("value_or_part", "")}' for x in plan.get("components", [])[:10])
        elif index == 4:
            value = "\n".join(map(str, plan.get("firmware_modules", [])))
        elif index == 5:
            value = json.dumps(plan.get("pcb", {}), indent=2)
        elif index == 6:
            value = "\n".join(map(str, plan.get("validation_steps", [])))
        else:
            value = plan.get("novelty", "")
        body.text = value
    prs.save(path)

def write_pdf(path: Path, plan: dict):
    from reportlab.platypus import SimpleDocTemplate, Paragraph
    from reportlab.lib.styles import getSampleStyleSheet
    styles = getSampleStyleSheet()
    story = [
        Paragraph(escape(plan["name"]), styles["Title"]),
        Paragraph(escape(plan.get("objective", "")), styles["BodyText"]),
    ]
    for item in plan.get("validation_steps", []):
        story.append(Paragraph(escape(str(item)), styles["BodyText"]))
    SimpleDocTemplate(str(path)).build(story)

def write_diagram(path: Path, plan: dict):
    blocks = plan.get("blocks", []) or [
        {"name": "Controller", "x": 60, "y": 100, "width": 180, "height": 70}
    ]
    positions = {}
    output = [
        '<svg xmlns="http://www.w3.org/2000/svg" width="1000" height="650">',
        '<rect width="100%" height="100%" fill="white"/>',
        f'<text x="30" y="35" font-size="24" font-family="sans-serif">{escape(plan["name"])}</text>'
    ]
    for block in blocks:
        x = float(block.get("x", 50)); y = float(block.get("y", 100))
        width = float(block.get("width", 160)); height = float(block.get("height", 70))
        name = str(block.get("name", "Block"))
        positions[name] = (x, y, width, height)
        output.append(
            f'<rect x="{x}" y="{y}" width="{width}" height="{height}" rx="10" fill="#eef" stroke="#223" stroke-width="2"/>'
            f'<text x="{x + width/2}" y="{y + height/2}" text-anchor="middle" font-size="16">{escape(name)}</text>'
        )
    for connection in plan.get("connections", []):
        first = positions.get(str(connection.get("from")))
        second = positions.get(str(connection.get("to")))
        if first and second:
            x1 = first[0] + first[2]/2; y1 = first[1] + first[3]/2
            x2 = second[0] + second[2]/2; y2 = second[1] + second[3]/2
            output.append(f'<line x1="{x1}" y1="{y1}" x2="{x2}" y2="{y2}" stroke="#333" stroke-width="2"/>')
    output.append("</svg>")
    path.write_text("".join(output), encoding="utf-8")

def write_pcb_draft(path: Path, plan: dict):
    pcb = plan.get("pcb", {})
    width = float(pcb.get("width_mm", 80)); height = float(pcb.get("height_mm", 50))
    lines = [
        '(kicad_pcb (version 20240108) (generator "darshan-jarvis")',
        '(general (thickness 1.6))',
        '(layers (0 "F.Cu" signal) (31 "B.Cu" signal) (36 "B.SilkS" user "b.silkscreen") (37 "F.SilkS" user "f.silkscreen") (44 "Edge.Cuts" user))',
        '(setup (pad_to_mask_clearance 0))',
        f'(gr_rect (start 0 0) (end {width} {height}) (stroke (width 0.5) (type default)) (fill none) (layer "Edge.Cuts"))'
    ]
    for index, component in enumerate(plan.get("components", [])[:12]):
        x = 5 + (index % 6) * min(12, width / 7)
        y = 8 + (index // 6) * 12
        name = escape(str(component.get("name", "COMP"))[:24])
        lines.append(
            f'(footprint "JARVIS:GENERIC_{index+1}" (layer "F.Cu") (at {x:.2f} {y:.2f}) '
            f'(property "Reference" "J{index+1}" (at 0 -2 0) (layer "F.SilkS")) '
            f'(property "Value" "{name}" (at 0 2 0) (layer "F.Fab")) '
            '(fp_rect (start -2 -1.5) (end 2 1.5) (stroke (width 0.25) (type default)) (fill none) (layer "F.SilkS")) '
            '(pad "1" thru_hole circle (at 0 0) (size 2.4 2.4) (drill 1) (layers "*.Cu" "*.Mask")))'
        )
    lines.append(")")
    path.write_text("\n".join(lines), encoding="utf-8")

async def create_project(request: str) -> dict:
    plan = await make_plan(request)
    folder = WORK_ROOT / slug(plan.get("name", "jarvis_project"))
    folder.mkdir(parents=True, exist_ok=True)
    write_docx(folder / "project_report.docx", plan)
    write_pptx(folder / "project_presentation.pptx", plan)
    write_pdf(folder / "project_report.pdf", plan)
    write_diagram(folder / "final_architecture.svg", plan)
    write_pcb_draft(folder / "pcb_draft.kicad_pcb", plan)
    bom_rows = []\n    for i, component in enumerate(plan.get("components", [])):\n        name = str(component.get("name", ""))\n        value = str(component.get("value_or_part", ""))\n        qty = component.get("qty", 1)\n        bom_rows.append(f"J{i+1},{name},{value},{qty}")\n    (folder / "BOM.csv").write_text("Reference,Component,Value,Quantity\\n" + "\\n".join(bom_rows), encoding="utf-8")\n    (folder / "project_plan.json").write_text(json.dumps(plan, indent=2, ensure_ascii=False), encoding="utf-8")
    (folder / "PCB_README.md").write_text(
        "# PCB draft\n\nAutomatically generated engineering draft. "
        "Verify footprints, pin mapping, ERC/DRC, power integrity and manufacturing outputs before fabrication.\n",
        encoding="utf-8"
    )
    return {
        "name": plan["name"],
        "folder": str(folder),
        "files": [str(item) for item in sorted(folder.iterdir())],
        "plan": plan,
    }
